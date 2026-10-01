#!/usr/bin/env python3
"""Merge requirement-register JSON files into one compliance matrix (xlsx + json).

Input rows come from shredding a solicitation in segments (see SKILL.md step 2). Each
row: id, lines, section, type, requirement, quote, response_location, owner, risk, note.

The matrix adds what makes it a working tool rather than a list:
  * status column (open / answered / exception / n-a) with a dropdown
  * "answered in" column the verifier fills with the document and heading that answers it
  * summary by type, owner and risk, and a check that ids are unique and required fields exist
  * anomaly sheet (contradictions and gaps found while shredding)

Usage:
  compliance_matrix.py REGISTER_DIR OUT_PREFIX [--title "RFP-XXXX"]
    reads REGISTER_DIR/shred-*.json and REGISTER_DIR/shred-*-anomalies.md
    writes OUT_PREFIX.xlsx and OUT_PREFIX.json
  compliance_matrix.py --rebuild MATRIX.json OUT_PREFIX [--anomalies-from REGISTER_DIR]
    regenerates the workbook from a matrix JSON that already carries statuses and notes
    (re-running the first form from the shred files would reset every status to open)
  compliance_matrix.py --verify MATRIX.json   # counts rows still 'open' (exit 1 if any)
Requires: openpyxl.
"""
from __future__ import annotations

import argparse
import collections
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

FIELDS = ["id", "lines", "section", "type", "requirement", "quote", "response_location", "owner", "risk", "note"]
REQUIRED = ["id", "requirement", "type", "owner"]
RISK_FILL = {"high": "F8D7DA", "medium": "FFF3CD", "low": "E2F0D9"}
THIN = Side(style="thin", color="BFC5CC")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def load_rows(reg_dir: Path) -> tuple[list[dict], list[str]]:
    rows, problems = [], []
    # register files are exactly shred-<n>.json; helper files (e.g. shred-2-property-table.json) are not rows
    for f in sorted(p for p in reg_dir.glob("shred-*.json") if re.fullmatch(r"shred-\d+\.json", p.name)):
        try:
            data = json.loads(f.read_text(encoding="utf8"))
        except json.JSONDecodeError as exc:
            problems.append(f"{f.name}: invalid JSON ({exc})")
            continue
        if isinstance(data, dict):  # tolerate {"rows":[...]}
            data = data.get("rows", [])
        for r in data:
            r["_file"] = f.name
            rows.append(r)
    return rows, problems


def check(rows: list[dict]) -> list[str]:
    problems, seen = [], collections.Counter(r.get("id") for r in rows)
    for rid, n in seen.items():
        if n > 1:
            problems.append(f"duplicate id {rid} x{n}")
    for r in rows:
        for k in REQUIRED:
            if not str(r.get(k, "")).strip():
                problems.append(f"{r.get('id', '?')}: missing {k}")
        if r.get("risk") not in (None, "", "low", "medium", "high"):
            problems.append(f"{r.get('id')}: risk '{r.get('risk')}' not low/medium/high")
    return problems


def anomalies(reg_dir: Path) -> list[tuple[str, str]]:
    out = []
    for f in sorted(reg_dir.glob("shred-*-anomalies.md")):
        for ln in f.read_text(encoding="utf8").splitlines():
            ln = ln.strip()
            if ln and not ln.startswith("#"):
                out.append((f.name.replace("-anomalies.md", ""), re.sub(r"^[-*\d.)\s]+", "", ln)))
    return out


def build(reg_dir: Path, prefix: Path, title: str) -> int:
    rows, problems = load_rows(reg_dir)
    return write_outputs(rows, problems, reg_dir, prefix, title)


def rebuild(matrix_json: Path, prefix: Path, title: str, reg_dir: Path | None) -> int:
    """Regenerate the workbook from a matrix JSON that already carries statuses.

    Re-running `build` from the shred files would reset every status to open, so once a
    verifier has entered statuses and notes the JSON is the source of truth. The anomalies
    sheet still comes from the shred folder when one is given.
    """
    rows = json.loads(matrix_json.read_text(encoding="utf8"))
    return write_outputs(rows, [], reg_dir, prefix, title)


def write_outputs(rows: list[dict], problems: list[str], reg_dir: Path | None, prefix: Path, title: str) -> int:
    problems += check(rows)
    for r in rows:
        r.setdefault("status", "open")
        r.setdefault("answered_in", "")
    prefix.parent.mkdir(parents=True, exist_ok=True)
    Path(f"{prefix}.json").write_text(json.dumps(rows, indent=1), encoding="utf8")

    wb = Workbook()
    ws = wb.active
    ws.title = "Matrix"
    heads = FIELDS + ["status", "answered_in", "status_note"]
    for c, h in enumerate(heads, 1):
        cell = ws.cell(row=1, column=c, value=h.replace("_", " ").title())
        cell.font, cell.fill, cell.border = Font(bold=True), PatternFill("solid", fgColor="E9EDF2"), BOX
    for i, r in enumerate(rows, 2):
        for c, h in enumerate(heads, 1):
            cell = ws.cell(row=i, column=c, value=str(r.get(h, "")) if r.get(h) is not None else "")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = BOX
        risk = r.get("risk")
        if risk in RISK_FILL:
            ws.cell(row=i, column=FIELDS.index("risk") + 1).fill = PatternFill("solid", fgColor=RISK_FILL[risk])
    widths = [10, 9, 24, 14, 60, 46, 26, 16, 8, 50, 10, 30, 50]
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(heads)).column_letter}{len(rows) + 1}"
    dv = DataValidation(type="list", formula1='"open,answered,exception,n-a"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"K2:K{len(rows) + 1}")

    sm = wb.create_sheet("Summary")
    sm["A1"] = f"{title} - requirement register"
    sm["A1"].font = Font(bold=True, size=13)
    r0 = 3
    for label, key in (("By type", "type"), ("By owner", "owner"), ("By risk", "risk")):
        sm.cell(row=r0, column=1, value=label).font = Font(bold=True)
        for j, (k, n) in enumerate(sorted(collections.Counter(r.get(key) or "(blank)" for r in rows).items()), r0 + 1):
            sm.cell(row=j, column=1, value=k)
            sm.cell(row=j, column=2, value=n)
        r0 += 3 + len(set(r.get(key) or "(blank)" for r in rows))
    sm.cell(row=r0, column=1, value="Total rows").font = Font(bold=True)
    sm.cell(row=r0, column=2, value=len(rows))
    sm.cell(row=r0 + 1, column=1, value="Still open").font = Font(bold=True)
    sm.cell(row=r0 + 1, column=2, value=f'=COUNTIF(Matrix!K2:K{len(rows) + 1},"open")')
    sm.column_dimensions["A"].width = 34

    an = wb.create_sheet("Anomalies")
    an["A1"], an["B1"] = "Segment", "Contradiction / gap / stale reference"
    an["A1"].font = an["B1"].font = Font(bold=True)
    anomaly_rows = anomalies(reg_dir) if reg_dir else []
    for i, (seg, text) in enumerate(anomaly_rows, 2):
        an.cell(row=i, column=1, value=seg)
        c = an.cell(row=i, column=2, value=text)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    an.column_dimensions["A"].width = 12
    an.column_dimensions["B"].width = 140
    wb.save(f"{prefix}.xlsx")

    n_files = len({r.get("_file") for r in rows})
    by_risk = collections.Counter(r.get("risk") for r in rows)
    by_status = collections.Counter(r.get("status") for r in rows)
    print(f"{prefix}.xlsx: {len(rows)} rows from {n_files} files; "
          f"risk high={by_risk.get('high', 0)} medium={by_risk.get('medium', 0)} low={by_risk.get('low', 0)}; "
          f"status {dict(sorted(by_status.items()))}; anomalies={len(anomaly_rows)}")
    for p in problems:
        print("PROBLEM:", p)
    return 1 if problems else 0


def verify(path: Path) -> int:
    rows = json.loads(path.read_text(encoding="utf8"))
    open_rows = [r["id"] for r in rows if r.get("status", "open") == "open"]
    print(f"{len(rows)} rows, {len(open_rows)} still open")
    if open_rows:
        print("open:", ", ".join(open_rows[:40]) + (" ..." if len(open_rows) > 40 else ""))
    return 1 if open_rows else 0


def main(argv: list[str]) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("a")
    ap.add_argument("b", nargs="?")
    ap.add_argument("--title", default="Solicitation")
    ap.add_argument("--verify", action="store_true")
    ap.add_argument("--rebuild", action="store_true")
    ap.add_argument("--anomalies-from", default=None)
    args = ap.parse_args(argv)
    if args.verify:
        return verify(Path(args.a))
    if args.rebuild:
        if not args.b:
            ap.error("--rebuild needs MATRIX.json and OUT_PREFIX")
        return rebuild(Path(args.a), Path(args.b), args.title, Path(args.anomalies_from) if args.anomalies_from else None)
    if not args.b:
        ap.error("need REGISTER_DIR and OUT_PREFIX")
    return build(Path(args.a), Path(args.b), args.title)


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
