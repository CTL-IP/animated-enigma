#!/usr/bin/env python3
"""Markdown subset -> clean Word document, with fact tokens resolved from an intake file.

Why this exists: a bid is many documents that must agree on the same facts. Drafts are
written once as Markdown with {{firm.legal_name}}-style tokens; this script resolves
them from the intake YAML. A token whose value is null/missing becomes a yellow
[[FILL: key]] flag instead of a guess, and `scan_placeholders.py` counts them.

Usage:
  md_to_docx.py draft.md out.docx [--intake intake.yaml] [--accent F76808]

Front matter (optional, YAML between --- lines at the top of the .md):
  title, subtitle, header_left, header_right, footer_note, cover (true/false)

Supported Markdown: # to #### headings, paragraphs, - / * / 1. lists (two nesting
levels), > callouts, GFM pipe tables, ``` code, --- rule, <!-- pagebreak -->,
**bold**, *italic*, `code`, [ ] / [x] checkboxes, and [[FILL: ...]] flags.

Requires: python-docx, pyyaml.
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

import yaml
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FILL_RE = re.compile(r"\[\[(.*?)\]\]")
TOKEN_RE = re.compile(r"\{\{\s*([A-Za-z0-9_.\[\]-]+)\s*\}\}")
INLINE_RE = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`.+?`|\[\[.*?\]\])")


# --------------------------------------------------------------------------- tokens
def load_intake(path: str | None) -> dict:
    if not path:
        return {}
    with open(path, encoding="utf8") as fh:
        return yaml.safe_load(fh) or {}


def lookup(data: dict, dotted: str):
    cur = data
    for part in dotted.split("."):
        m = re.fullmatch(r"([A-Za-z0-9_-]+)\[(\d+)\]", part)
        if m:
            cur = cur.get(m.group(1)) if isinstance(cur, dict) else None
            try:
                cur = cur[int(m.group(2))] if cur is not None else None
            except (IndexError, TypeError):
                return None
        else:
            cur = cur.get(part) if isinstance(cur, dict) else None
        if cur is None:
            return None
    return cur


def resolve_tokens(text: str, intake: dict) -> tuple[str, list[str]]:
    """Replace {{a.b}} with the value, or with [[FILL: a.b]] when unknown."""
    missing: list[str] = []

    def repl(m: re.Match) -> str:
        val = lookup(intake, m.group(1))
        if val is None or val == "" or isinstance(val, (dict, list)):
            missing.append(m.group(1))
            return f"[[FILL: {m.group(1)}]]"
        return str(val)

    return TOKEN_RE.sub(repl, text), missing


def split_front_matter(text: str) -> tuple[dict, str]:
    if text.startswith("---\n"):
        end = text.find("\n---", 4)
        if end != -1:
            meta = yaml.safe_load(text[4:end]) or {}
            return meta, text[end + 4:].lstrip("\n")
    return {}, text


# ---------------------------------------------------------------------- docx helpers
def shade(cell_or_par, hex_fill: str) -> None:
    el = cell_or_par._tc if hasattr(cell_or_par, "_tc") else cell_or_par._p
    pr = el.get_or_add_tcPr() if hasattr(el, "get_or_add_tcPr") else el.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    pr.append(shd)


def add_field(run, instr: str) -> None:
    for kind, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        elif text == instr:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = instr
        else:
            el = OxmlElement("w:t")
            el.text = text
        run._r.append(el)


def add_runs(par, text: str, size: float | None = None, base_bold=False, color: RGBColor | None = None) -> None:
    for chunk in INLINE_RE.split(text):
        if not chunk:
            continue
        run_text, bold, italic, code, flag = chunk, base_bold, False, False, False
        if chunk.startswith("**") and chunk.endswith("**") and len(chunk) > 4:
            run_text, bold = chunk[2:-2], True
        elif chunk.startswith("`") and chunk.endswith("`") and len(chunk) > 2:
            run_text, code = chunk[1:-1], True
        elif chunk.startswith("[[") and chunk.endswith("]]"):
            run_text, flag, bold = chunk, True, True
        elif chunk.startswith("*") and chunk.endswith("*") and len(chunk) > 2:
            run_text, italic = chunk[1:-1], True
        run_text = run_text.replace("[ ]", "☐").replace("[x]", "☒").replace("[X]", "☒")
        run = par.add_run(run_text)
        run.bold = bold
        run.italic = italic
        if code:
            run.font.name = "Consolas"
        if flag:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        if size:
            run.font.size = Pt(size)
        if color is not None and not flag:
            run.font.color.rgb = color


def set_repeat_header(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def no_split(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def setup_document(meta: dict, accent: str) -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.orientation = WD_ORIENT.PORTRAIT
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Inches(1))
    sec.top_margin, sec.bottom_margin = Inches(1), Inches(0.9)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for name in ("Header", "Footer"):
        doc.styles[name].font.size = Pt(8.5)
        doc.styles[name].font.name = "Calibri"
    ink = RGBColor(0x1F, 0x2A, 0x37)
    sizes = {1: 20, 2: 15, 3: 12.5, 4: 11}
    for lvl, size in sizes.items():
        st = doc.styles[f"Heading {lvl}"]
        st.font.name = "Calibri"
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = ink
        st.paragraph_format.space_before = Pt(16 if lvl <= 2 else 10)
        st.paragraph_format.space_after = Pt(6 if lvl <= 2 else 4)
        st.paragraph_format.keep_with_next = True
    # accent rule under H1
    h1 = doc.styles["Heading 1"]
    pPr = h1.element.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "12"), ("space", "2"), ("color", accent.upper())):
        bottom.set(qn(f"w:{k}"), v)
    bdr.append(bottom)
    pPr.append(bdr)

    # header / footer
    hdr = sec.header.paragraphs[0]
    hdr.text = ""
    left = meta.get("header_left", "")
    right = meta.get("header_right", "")
    if left or right:
        add_runs(hdr, left, size=8.5, color=RGBColor(0x55, 0x60, 0x6E))
        hdr.add_run("\t\t")
        add_runs(hdr, right, size=8.5, color=RGBColor(0x55, 0x60, 0x6E))
        tabs = hdr.paragraph_format.tab_stops
        tabs.add_tab_stop(Inches(6.5), alignment=2)  # right
    ftr = sec.footer.paragraphs[0]
    ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note = meta.get("footer_note", "")
    if note:
        r = ftr.add_run(note + "   |   ")
        r.font.size = Pt(8.5)
    r = ftr.add_run("Page ")
    r.font.size = Pt(8.5)
    r2 = ftr.add_run()
    r2.font.size = Pt(8.5)
    add_field(r2, "PAGE")
    r3 = ftr.add_run(" of ")
    r3.font.size = Pt(8.5)
    r4 = ftr.add_run()
    r4.font.size = Pt(8.5)
    add_field(r4, "NUMPAGES")
    return doc


def add_cover(doc: Document, meta: dict, accent: str) -> None:
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    run = p.add_run(meta.get("title", ""))
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor(0x1F, 0x2A, 0x37)
    if meta.get("subtitle"):
        p2 = doc.add_paragraph()
        add_runs(p2, meta["subtitle"], size=14, color=RGBColor(0x55, 0x60, 0x6E))
    bar = doc.add_paragraph()
    bar.paragraph_format.space_before = Pt(12)
    pPr = bar._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "24"), ("space", "1"), ("color", accent.upper())):
        b.set(qn(f"w:{k}"), v)
    bdr.append(b)
    pPr.append(bdr)
    for line in meta.get("cover_lines", []):
        p = doc.add_paragraph()
        add_runs(p, str(line), size=11.5)
    doc.add_page_break()


# ------------------------------------------------------------------------ block parse
def parse_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for ln in lines:
        ln = ln.strip()
        if re.fullmatch(r"\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?", ln):
            continue
        cells = [c.strip() for c in re.split(r"(?<!\\)\|", ln.strip("|"))]
        rows.append([c.replace("\\|", "|") for c in cells])
    return rows


def add_table(doc: Document, rows: list[list[str]], accent: str) -> None:
    ncols = max(len(r) for r in rows)
    tbl = doc.add_table(rows=len(rows), cols=ncols)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, r in enumerate(rows):
        row = tbl.rows[i]
        no_split(row)
        for j in range(ncols):
            cell = row.cells[j]
            cell.text = ""
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(2)
            txt = r[j] if j < len(r) else ""
            add_runs(par, txt.replace("<br>", "\n"), size=9.5, base_bold=(i == 0))
            if i == 0:
                shade(cell, "E9EDF2")
        if i == 0:
            set_repeat_header(row)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def convert(text: str, intake: dict, accent: str) -> tuple[Document, list[str]]:
    meta, body = split_front_matter(text)
    meta_text, missing_meta = resolve_tokens(yaml.safe_dump(meta, allow_unicode=True), intake)
    meta = yaml.safe_load(meta_text) or {}
    body, missing = resolve_tokens(body, intake)
    doc = setup_document(meta, accent)
    if meta.get("cover"):
        add_cover(doc, meta, accent)
    lines = body.split("\n")
    i = 0
    para_buf: list[str] = []

    def flush_para() -> None:
        nonlocal para_buf
        if para_buf:
            p = doc.add_paragraph()
            add_runs(p, " ".join(s.strip() for s in para_buf))
            para_buf = []

    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if not s:
            flush_para(); i += 1; continue
        if s.startswith("<!--") and "pagebreak" in s:
            flush_para(); doc.add_page_break(); i += 1; continue
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            flush_para()
            h = doc.add_heading(level=len(m.group(1)))
            add_runs(h, m.group(2))
            i += 1; continue
        if s.startswith("```"):
            flush_para()
            i += 1
            code = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i]); i += 1
            i += 1
            p = doc.add_paragraph()
            shade(p, "F3F4F6")
            r = p.add_run("\n".join(code))
            r.font.name = "Consolas"; r.font.size = Pt(9)
            continue
        if re.fullmatch(r"-{3,}|\*{3,}", s):
            flush_para()
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            bdr = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
            for k, v in (("val", "single"), ("sz", "4"), ("space", "1"), ("color", "BFC5CC")):
                b.set(qn(f"w:{k}"), v)
            bdr.append(b); pPr.append(bdr)
            i += 1; continue
        if s.startswith("|"):
            flush_para()
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i]); i += 1
            rows = parse_table(block)
            if rows:
                add_table(doc, rows, accent)
            continue
        if s.startswith(">"):
            flush_para()
            block = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                block.append(lines[i].strip().lstrip(">").strip()); i += 1
            p = doc.add_paragraph()
            shade(p, "FFF4E8")
            p.paragraph_format.left_indent = Inches(0.15)
            add_runs(p, " ".join(block))
            continue
        lm = re.match(r"^(\s*)([-*]|\d+[.)])\s+(.*)$", ln)
        if lm:
            flush_para()
            indent = len(lm.group(1).replace("\t", "    "))
            level = 1 if indent >= 2 else 0
            ordered = lm.group(2)[0].isdigit()
            style = ("List Number" if ordered else "List Bullet") + (" 2" if level else "")
            p = doc.add_paragraph(style=style)
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, lm.group(3))
            i += 1; continue
        para_buf.append(ln)
        i += 1
    flush_para()
    return doc, missing + missing_meta


def main(argv: list[str]) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("src")
    ap.add_argument("out")
    ap.add_argument("--intake")
    ap.add_argument("--accent", default="F76808")
    args = ap.parse_args(argv)
    text = Path(args.src).read_text(encoding="utf8")
    doc, missing = convert(text, load_intake(args.intake), args.accent)
    Path(args.out).parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.out)
    flags = len(FILL_RE.findall(text)) + len(missing)
    print(f"{args.out}: written; unresolved fact tokens={len(missing)}; total [[FILL]] flags in source+tokens={flags}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
