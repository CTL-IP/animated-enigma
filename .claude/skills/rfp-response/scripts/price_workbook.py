#!/usr/bin/env python3
"""Build a unit-price bid workbook from a solicitation's price schedule.

The rule this encodes: a price is computed from the firm's OWN inputs or it is blank.
Nothing is pre-filled, and a blank rate can never silently price as $0 -- every computed
cell checks ISNUMBER on every input it needs and shows "" until they are all real.

Two ways to get a schedule.json:
  price_workbook.py parse RFP_TEXT.txt FIRST-LAST OUT.json
      Parse "Item | UOM | $ -" rows under "Section | UOM | Price" headers
      (RFP_TEXT.txt comes from extract_rfp.py; lines are 1-indexed).
  ... or write it by hand / with an agent (see SCHEDULE FORMAT).
Then:
  price_workbook.py build schedule.json OUT.xlsx [--benchmarks price-points.csv]

SCHEDULE FORMAT
{
  "title": "Unit Price Schedule - RFP-XXXX",
  "sections": [ {"name": "Painting - Sections (Labor + Materials)", "trade": "painter",
                 "paint_materials": true,
                 "items": [ {"id": "A-007-003", "item": "Accent Wall", "uom": "sf",
                             "basis": "L+M", "trade": "painter(optional override)",
                             "hours": 0.02, "hours_confidence": "M", "hours_note": "why this many hours",
                             "needs_materials": true, "sub_required": false} ]} ],
  "summary": [ {"name": "Carpet Cleaning (Labor Only)",
                "items": [ {"item": "1 Bedroom", "uom": "each", "ref": "A-004-001"},
                           {"item": "B-only line", "uom": "each", "ref": null} ]} ],
  "site_table": {"sizes": ["Efficiency","1 Bdrm","2 Bdrm","3 Bdrm","4 Bdrm","5 Bdrm"],
                 "properties": [ {"name": "...", "units": 121, "address": "...",
                                  "shaded_cells": ["Efficiency", "3 Bdrm"]} ]}
  shaded_cells (optional, per property): sizes the agency greyed out. They are drawn grey, are NOT input
  cells, and do not count toward READY. Without the key every cell is an input.
}
Item fields (all optional): hours (starting ASSUMPTION per unit), hours_confidence (H/M/L), hours_note,
needs_materials (price stays blank until a real material $/unit is entered), sub_required (price comes from
a subcontract quote, not hours x rate).
Inputs can be prefilled from a JSON file:  build ... --inputs inputs.json
  {"overhead": 0.12, "profit": 0.15, "tax": 0.0825, "dha_paint": "N", "round_to": 0.05,
   "rates": {"cleaner": 51.2, "painter": 51.2}}
Only numbers that came from the owner belong there.
Workbook sheets: README, Inputs, Schedule A, Attachment B, Site Table, Checks, Benchmarks.
Requires: openpyxl.
"""
from __future__ import annotations

import csv
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

TRADES = ["cleaner", "painter", "drywall", "carpenter", "electrician", "plumber", "flooring", "general_laborer"]
SECTION_TRADE = [
    ("clean", "cleaner"), ("carpet", "cleaner"), ("trash", "general_laborer"), ("additional", "cleaner"),
    ("paint", "painter"), ("finish", "drywall"), ("drywall", "drywall"), ("electric", "electrician"),
    ("plumb", "plumber"), ("floor", "flooring"), ("cabinet", "carpenter"), ("door", "carpenter"),
    ("window", "carpenter"), ("fence", "carpenter"),
]
INPUT_FILL = PatternFill("solid", fgColor="FFF8DC")
HEAD_FILL = PatternFill("solid", fgColor="E9EDF2")
BLUE = Font(color="0000FF")
BOLD = Font(bold=True)
THIN = Side(style="thin", color="BFC5CC")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def guess_trade(section: str) -> str:
    low = section.lower()
    for key, trade in SECTION_TRADE:
        if key in low:
            return trade
    return ""


# ------------------------------------------------------------------------------ parse
def parse(text_path: str, span: str, out: str) -> None:
    first, last = (int(x) for x in span.split("-"))
    lines = Path(text_path).read_text(encoding="utf8").split("\n")[first - 1:last]
    sections, cur, n_sec = [], None, 0
    for ln in lines:
        cells = [c.strip() for c in ln.split(" | ")]
        if len(cells) == 3 and cells[1].lower() in {"uom", "unit"} and cells[2].lower() == "price":
            n_sec += 1
            name = cells[0]
            cur = {"name": name, "trade": guess_trade(name),
                   "paint_materials": "paint" in name.lower() and "material" in name.lower().replace("mat.", "material"),
                   "items": []}
            sections.append(cur)
        elif cur is not None and len(cells) == 3 and re.fullmatch(r"\$\s*-?", cells[2]):
            cur["items"].append({"id": f"A-{n_sec:02d}-{len(cur['items']) + 1:03d}", "item": cells[0], "uom": cells[1],
                                 "basis": "L+M" if re.search(r"\(L\+M\)|Labor\s*[+&]\s*Mat", cells[0] + " " + cur["name"], re.I) else "L"})
    Path(out).write_text(json.dumps({"title": "Unit Price Schedule", "sections": sections}, indent=2), encoding="utf8")
    print(f"{out}: {len(sections)} sections, {sum(len(s['items']) for s in sections)} priced items")


# ------------------------------------------------------------------------------ build
def style_header(ws, row: int, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font, cell.fill, cell.border = BOLD, HEAD_FILL, BOX
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def build(schedule_path: str, out: str, bench_csv: str | None, inputs: dict | None = None) -> None:
    sch = json.loads(Path(schedule_path).read_text(encoding="utf8"))
    wb = Workbook()

    # README
    ws = wb.active
    ws.title = "README"
    for i, line in enumerate([
        sch.get("title", "Unit Price Schedule"),
        "",
        "How this workbook works",
        "1. Fill the yellow cells on Inputs: burdened hourly rate for each trade, overhead %, profit %, sales-tax rate on materials, and whether DHA furnishes paint.",
        "2. On Schedule A, enter hours per unit (and material / subcontract cost per unit where the line includes them) in the yellow cells. Blue text = your input; black = formula.",
        "3. A price appears only when EVERY input it needs is a real number. A blank rate does not price as $0 -- the cell stays empty.",
        "4. Override price (column K) wins over the computed price: use it when you price a line by judgment. Explain overrides in Notes.",
        "5. Attachment B pulls its prices from Schedule A through the Ref ID, so the two schedules cannot disagree. B-only lines take their own input.",
        "6. Hours per unit are STARTING ASSUMPTIONS (estimator judgement, with a confidence letter and a note). They are not your production data: review every one, change what your crews would really do, then set 'reviewed' on Inputs to Y.",
        "7. A line marked Materials needed stays blank until you enter a real material $/unit (price it at the store, with the sales-tax decision made). A line marked Sub required stays blank until you enter a subcontractor quote.",
        "8. Checks shows what is still blank and whether the package is READY. Do not upload until it says READY.",
        "",
        "Nothing in this workbook is a price until you enter your inputs. Starting hours, where present, are assumptions to verify against your crews' actual production.",
    ], 1):
        ws.cell(row=i, column=1, value=line)
    ws["A1"].font = Font(bold=True, size=14)
    ws["A3"].font = BOLD
    ws.column_dimensions["A"].width = 140

    # Inputs
    wi = wb.create_sheet("Inputs")
    wi["A1"], wi["A1"].font = "Pricing inputs (yellow = you fill)", Font(bold=True, size=12)
    rows = [("Overhead %  (e.g. 0.12 for 12%)", "OVH"), ("Profit %  (e.g. 0.10 for 10%)", "PROFIT"),
            ("Sales tax on materials you buy  (see tax research; e.g. 0.0825)", "TAX"),
            ("DHA furnishes paint?  (Y / N)", "DHA_PAINT"), ("Round prices to nearest $  (e.g. 0.25)", "ROUNDTO"),
            ("Owner has reviewed every starting assumption (hours, trades, basis)?  (Y / N)", "REVIEWED")]
    for i, (label, name) in enumerate(rows, 3):
        wi.cell(row=i, column=1, value=label)
        c = wi.cell(row=i, column=2)
        c.fill, c.font, c.border = INPUT_FILL, BLUE, BOX
        wb.defined_names[name] = DefinedName(name, attr_text=f"Inputs!$B${i}")
    dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    wi.add_data_validation(dv)
    dv.add("B6")
    dv.add("B8")
    wi["A11"], wi["B11"] = "Trade", "Burdened labor rate $/hour"
    style_header(wi, 11, 2)
    for i, t in enumerate(TRADES, 12):
        wi.cell(row=i, column=1, value=t)
        c = wi.cell(row=i, column=2)
        c.fill, c.font, c.border = INPUT_FILL, BLUE, BOX
        c.number_format = '"$"#,##0.00'
    last_rate = 11 + len(TRADES)
    wb.defined_names["RK"] = DefinedName("RK", attr_text=f"Inputs!$A$12:$A${last_rate}")
    wb.defined_names["RV"] = DefinedName("RV", attr_text=f"Inputs!$B$12:$B${last_rate}")
    if inputs:
        for cell, key in (("B3", "overhead"), ("B4", "profit"), ("B5", "tax"), ("B6", "dha_paint"), ("B7", "round_to")):
            if inputs.get(key) is not None:
                wi[cell] = inputs[key]
        for i in range(12, last_rate + 1):
            v = (inputs.get("rates") or {}).get(wi.cell(row=i, column=1).value)
            if v is not None:
                wi.cell(row=i, column=2, value=v)
    wi.cell(row=last_rate + 2, column=1,
            value="Burdened = wages + payroll taxes + workers' comp + insurance + vehicle/tools per productive hour. Not the wage.")
    wi.column_dimensions["A"].width = 62
    wi.column_dimensions["B"].width = 26

    # Schedule A
    sa = wb.create_sheet("Schedule A")
    heads = ["ID", "Section", "Item", "UOM", "Basis", "Trade", "Hours / unit", "Material $ / unit", "Sub $ / unit",
             "Computed price", "Override price", "FINAL price", "Benchmark low", "Benchmark high", "Flag", "Notes",
             "Paint material?", "Materials needed?", "Sub required?", "Hours confidence"]
    for c, h in enumerate(heads, 1):
        sa.cell(row=1, column=c, value=h)
    style_header(sa, 1, len(heads))
    r = 2
    for sec in sch["sections"]:
        for it in sec["items"]:
            trade = it.get("trade") or sec.get("trade") or guess_trade(sec["name"])
            pm = "Y" if (sec.get("paint_materials") or it.get("paint_materials")) else "N"
            vals = [it["id"], sec["name"], it["item"], it["uom"], it.get("basis", ""), trade]
            for c, v in enumerate(vals, 1):
                sa.cell(row=r, column=c, value=v).border = BOX
            sa.cell(row=r, column=7, value=it.get("hours"))
            sa.cell(row=r, column=16, value=it.get("hours_note", ""))
            sa.cell(row=r, column=17, value=pm)
            sa.cell(row=r, column=18, value="Y" if it.get("needs_materials") else "N")
            sa.cell(row=r, column=19, value="Y" if it.get("sub_required") else "N")
            sa.cell(row=r, column=20, value=it.get("hours_confidence", ""))
            for c in (7, 8, 9, 11, 13, 14, 16):
                cell = sa.cell(row=r, column=c)
                cell.fill, cell.font, cell.border = INPUT_FILL, BLUE, BOX
            mat = f'IF(AND(DHA_PAINT="Y",Q{r}="Y"),0,N(H{r}))'
            needmat = f'AND(R{r}="Y",NOT(AND(DHA_PAINT="Y",Q{r}="Y")))'
            rate = f'INDEX(RV,MATCH(F{r},RK,0))'
            labor = f'IF(S{r}="Y",0,G{r}*{rate})'
            base = f'({labor}+{mat}*(1+N(TAX))+N(I{r}))*(1+OVH)*(1+PROFIT)'
            blocked = (
                f'OR(NOT(ISNUMBER(OVH)),NOT(ISNUMBER(PROFIT)),'
                f'AND(S{r}="Y",NOT(ISNUMBER(I{r}))),'
                f'AND(S{r}<>"Y",OR(F{r}="",NOT(ISNUMBER(G{r})),NOT(ISNUMBER({rate})))),'
                f'AND({needmat},NOT(ISNUMBER(H{r}))),'
                f'AND({mat}>0,NOT(ISNUMBER(TAX))))'
            )
            computed = (
                f'=IFERROR(IF({blocked},"",IF(ISNUMBER(ROUNDTO),MROUND({base},ROUNDTO),ROUND({base},2))),"")'
            )
            sa.cell(row=r, column=10, value=computed).border = BOX
            sa.cell(row=r, column=12, value=f'=IF(ISNUMBER(K{r}),K{r},J{r})').border = BOX
            sa.cell(row=r, column=15, value=(
                f'=IF(NOT(ISNUMBER(L{r})),"",IF(AND(ISNUMBER(M{r}),L{r}<M{r}),"below benchmark",'
                f'IF(AND(ISNUMBER(N{r}),L{r}>N{r}),"above benchmark","")))')).border = BOX
            for c in (8, 9, 10, 11, 12, 13, 14):
                sa.cell(row=r, column=c).number_format = '"$"#,##0.00'
            r += 1
    last_a = r - 1
    for col, w in zip("ABCDEFGHIJKLMNOPQRST", [11, 34, 46, 9, 7, 14, 10, 12, 10, 12, 12, 12, 11, 11, 16, 40, 9, 10, 9, 10]):
        sa.column_dimensions[col].width = w
    sa.freeze_panes = "D2"
    sa.auto_filter.ref = f"A1:T{last_a}"

    # Attachment B
    sb = wb.create_sheet("Attachment B")
    for c, h in enumerate(["Section", "Item", "UOM", "Ref ID (Schedule A)", "B-only price (input)", "PRICE"], 1):
        sb.cell(row=1, column=c, value=h)
    style_header(sb, 1, 6)
    rb = 2
    for sec in sch.get("summary", []):
        for it in sec["items"]:
            for c, v in enumerate([sec["name"], it["item"], it["uom"], it.get("ref") or ""], 1):
                sb.cell(row=rb, column=c, value=v).border = BOX
            e = sb.cell(row=rb, column=5)
            e.fill, e.font, e.border = INPUT_FILL, BLUE, BOX
            sb.cell(row=rb, column=6, value=(
                f'=IF(D{rb}<>"",IFERROR(INDEX(\'Schedule A\'!$L$2:$L${last_a},MATCH(D{rb},\'Schedule A\'!$A$2:$A${last_a},0)),""),'
                f'IF(ISNUMBER(E{rb}),E{rb},""))')).border = BOX
            sb.cell(row=rb, column=6).number_format = '"$"#,##0.00'
            e.number_format = '"$"#,##0.00'
            rb += 1
    last_b = rb - 1
    for col, w in zip("ABCDEF", [36, 44, 10, 18, 18, 14]):
        sb.column_dimensions[col].width = w
    sb.freeze_panes = "C2"

    # Site table
    st = wb.create_sheet("Site Table")
    site = sch.get("site_table") or {"sizes": [], "properties": []}
    sizes = site["sizes"]
    GREY = PatternFill("solid", fgColor="D9D9D9")
    helper_col = 5 + len(sizes)
    for c, h in enumerate(["#", "Property", "Units", "Address"] + sizes + ["Open cells filled"], 1):
        st.cell(row=1, column=c, value=h)
    style_header(st, 1, helper_col)
    n_open_total = 0
    for i, p in enumerate(site["properties"], 2):
        for c, v in enumerate([i - 1, p.get("name") or "", p.get("units"), p.get("address") or ""], 1):
            st.cell(row=i, column=c, value=v).border = BOX
        shaded = set(p.get("shaded_cells") or [])
        open_refs = []
        for k, size in enumerate(sizes):
            col = 5 + k
            cell = st.cell(row=i, column=col)
            cell.border = BOX
            if size in shaded:
                cell.fill = GREY
            else:
                cell.fill, cell.font = INPUT_FILL, BLUE
                cell.number_format = '"$"#,##0.00'
                open_refs.append(f"{get_column_letter(col)}{i}")
        n_open_total += len(open_refs)
        st.cell(row=i, column=helper_col, value=f"=COUNT({','.join(open_refs)})" if open_refs else 0).border = BOX
    last_s = 1 + len(site["properties"])
    st.column_dimensions["B"].width = 34
    st.column_dimensions["D"].width = 38
    st.freeze_panes = "C2"
    st.cell(row=last_s + 2, column=2, value="Grey = shaded by the agency (meaning not defined: ask). Grey cells are not inputs and do not count toward READY.")

    # Benchmarks
    bm = wb.create_sheet("Benchmarks")
    if bench_csv and Path(bench_csv).exists():
        with open(bench_csv, newline="", encoding="utf8") as fh:
            for i, row in enumerate(csv.reader(fh), 1):
                for c, v in enumerate(row, 1):
                    bm.cell(row=i, column=c, value=v)
        style_header(bm, 1, bm.max_column)
    else:
        bm["A1"] = "No benchmark file supplied. Only cited public figures belong here."

    # Checks
    ck = wb.create_sheet("Checks")
    n_site_cells = n_open_total
    checks = [
        ("Schedule A lines", f"=COUNTA('Schedule A'!A2:A{last_a})"),
        ("Schedule A lines priced", f"=COUNT('Schedule A'!L2:L{last_a})"),
        ("Schedule A lines still blank", "=B3-B4"),
        ("Attachment B lines", f"=COUNTA('Attachment B'!B2:B{last_b})"),
        ("Attachment B lines priced", f"=COUNT('Attachment B'!F2:F{last_b})"),
        ("Attachment B lines still blank", "=B6-B7"),
        ("Site-table price cells expected (open, not grey)", n_site_cells),
        ("Site-table price cells filled", f"=SUM('Site Table'!{get_column_letter(helper_col)}2:{get_column_letter(helper_col)}{max(last_s, 2)})"),
        ("Site-table cells still blank", "=B9-B10"),
        ("Prices <= 0 (never valid)", f"=COUNTIF('Schedule A'!L2:L{last_a},\"<=0\")+COUNTIF('Attachment B'!F2:F{last_b},\"<=0\")"),
        ("Lines outside benchmark band", f"=COUNTIF('Schedule A'!O2:O{last_a},\"*benchmark\")"),
        ("Inputs complete? (overhead, profit)", '=IF(AND(ISNUMBER(OVH),ISNUMBER(PROFIT)),"overhead/profit set","overhead/profit MISSING")'),
        ("Owner reviewed every starting assumption?", '=IF(REVIEWED="Y","yes","NO - not reviewed")'),
        ("READY TO SUBMIT?", '=IF(AND(B5=0,B8=0,B11=0,B12=0,REVIEWED="Y"),"READY","NOT READY")'),
    ]
    ck["A1"], ck["B1"] = "Check", "Result"
    style_header(ck, 1, 2)
    for i, (label, formula) in enumerate(checks, 3):  # first check sits on row 3; formulas above refer to B3..B15
        ck.cell(row=i, column=1, value=label)
        ck.cell(row=i, column=2, value=formula)
    ck.column_dimensions["A"].width = 58
    ck.column_dimensions["B"].width = 26
    ck.cell(row=len(checks) + 4, column=1, value="Price cells in this workbook").font = BOLD
    ck.cell(row=len(checks) + 4, column=2, value="=B3+B6+B9")
    wb.save(out)
    print(f"{out}: Schedule A {last_a - 1} lines, Attachment B {last_b - 1} lines, site table {n_site_cells} cells "
          f"({last_a - 1 + last_b - 1 + n_site_cells} price points)")


def main(argv: list[str]) -> int:
    if len(argv) >= 4 and argv[0] == "parse":
        parse(argv[1], argv[2], argv[3])
        return 0
    if len(argv) >= 3 and argv[0] == "build":
        bench = argv[argv.index("--benchmarks") + 1] if "--benchmarks" in argv else None
        inp = json.loads(Path(argv[argv.index("--inputs") + 1]).read_text(encoding="utf8")) if "--inputs" in argv else None
        build(argv[1], argv[2], bench, inp)
        return 0
    print(__doc__)
    return 2


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
