#!/usr/bin/env python3
"""Self-test for the rfp-response scripts. Builds a synthetic solicitation and runs every
script the way a user would, asserting on real outputs. No company data, no network.

  python selftest.py            # prints each check and a final count; exit 1 on any failure
  python selftest.py --keep DIR # keep the working files in DIR for inspection

The fixture deliberately contains the traps that have bitten real solicitations:
  * text wrapped in <w:smartTag> (a naive reader drops it)
  * text inside a text box
  * a header naming a different solicitation number
  * a contents-list entry with no body text behind it
  * a price table in "Item | UOM | $ -" form under "Section | UOM | Price" headers
Formula recalculation needs LibreOffice (soffice); it is skipped, and said so, if absent.
Requires: python-docx, openpyxl, pyyaml (PyMuPDF and soffice optional).
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
PY = sys.executable
passed, failed = 0, []


def check(name: str, cond: bool, detail: str = "") -> None:
    global passed
    if cond:
        passed += 1
        print(f"  ok    {name}")
    else:
        failed.append(name)
        print(f"  FAIL  {name} {detail}")


def run(*args, expect_rc: int | None = 0) -> subprocess.CompletedProcess:
    cp = subprocess.run([PY, *map(str, args)], capture_output=True, text=True)
    if expect_rc is not None and cp.returncode != expect_rc:
        print(cp.stdout, cp.stderr)
    return cp


def smart_tag_paragraph(document, before: str, tagged: str, after: str):
    p = document.add_paragraph()
    p.add_run(before)
    tag = OxmlElement("w:smartTag")
    tag.set(qn("w:uri"), "urn:schemas-microsoft-com:office:smarttags")
    tag.set(qn("w:element"), "State")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = tagged
    r.append(t)
    tag.append(r)
    p._p.append(tag)
    p.add_run(after)
    return p


def textbox_paragraph(document, text: str):
    p = document.add_paragraph()
    r = OxmlElement("w:r")
    box = OxmlElement("w:txbxContent")
    bp = OxmlElement("w:p")
    br = OxmlElement("w:r")
    bt = OxmlElement("w:t")
    bt.text = text
    br.append(bt)
    bp.append(br)
    box.append(bp)
    r.append(box)
    p._p.append(r)


def build_fixture(path: Path) -> None:
    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = "Pursuant to: RFP-2024-52 SAMPLE CONTRACT"
    d.add_paragraph("REQUEST FOR PROPOSALS RFP-2026-99")
    d.add_paragraph("Proposals are due October 20, 2026 at 2:00 p.m.")
    d.add_paragraph("TABLE OF CONTENTS")
    for line in ("Introduction", "Scope of Services", "Certification of Compliance with Government Code 2271"):
        d.add_paragraph(line)
    d.add_paragraph("I. INTRODUCTION " + "The agency seeks qualified contractors to provide services under this solicitation. " * 3)
    smart_tag_paragraph(d, "This Contract is governed by the laws of the State of ", "Texas", ", venue in Dallas County.")
    textbox_paragraph(d, "BOXED CERTIFICATION TEXT")
    d.add_paragraph("Scope of Services: the contractor shall furnish labor and equipment.")
    t = d.add_table(rows=5, cols=3)
    rows = [("Cleaning (Labor Only)", "UOM", "Price"), ("1 bedroom unit", "each", "$ -"),
            ("Replace faucet (L+M)", "each", "$ -"), ("Replace toilet (L+M)", "each", "$ -"),
            ("Extra dirty oven", "each", "$ -")]
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            t.cell(i, j).text = v
    d.save(path)


def main(argv: list[str]) -> int:
    keep = Path(argv[argv.index("--keep") + 1]) if "--keep" in argv else None
    work = keep or Path(tempfile.mkdtemp(prefix="rfp-selftest-"))
    work.mkdir(parents=True, exist_ok=True)
    print(f"working in {work}")
    try:
        # ---------------------------------------------------------------- extract
        print("extract_rfp.py")
        src = work / "fixture.docx"
        build_fixture(src)
        out = work / "extract"
        cp = run(HERE / "extract_rfp.py", src, out)
        check("extractor exits 0", cp.returncode == 0, cp.stderr[-200:])
        txt = (out / "rfp.txt").read_text(encoding="utf8")
        inv = json.loads((out / "inventory.json").read_text(encoding="utf8"))
        check("smart-tagged text is kept (State of Texas)", "laws of the State of Texas, venue in Dallas County" in txt)
        check("text-box text appears exactly once", txt.count("BOXED CERTIFICATION TEXT") == 1)
        check("table rows use ' | ' cells", "1 bedroom unit | each | $ -" in txt)
        check("date is captured", "October 20, 2026" in inv["dates"])
        check("time is captured", any("2:00" in t for t in inv["times"]))
        check("foreign solicitation number in header is flagged",
              inv["other_solicitation_numbers_in_headers_or_footers"] == ["RFP-2024-52"],
              str(inv["other_solicitation_numbers_in_headers_or_footers"]))
        flagged = [e["entry"] for e in inv["contents_entries_without_body_text"]]
        check("contents entry with no body text is flagged, keeping its trailing number",
              any(e.endswith("Code 2271") for e in flagged), str(flagged))

        # ------------------------------------------------------- price workbook
        print("price_workbook.py")
        lines = txt.split("\n")
        start = next(i for i, l in enumerate(lines, 1) if l.startswith("Cleaning (Labor Only)"))
        sched = work / "schedule.json"
        cp = run(HERE / "price_workbook.py", "parse", out / "rfp.txt", f"{start}-{start + 4}", sched)
        data = json.loads(sched.read_text(encoding="utf8"))
        items = [it for s in data["sections"] for it in s["items"]]
        check("parse finds 1 section and 4 priced lines", len(data["sections"]) == 1 and len(items) == 4, cp.stdout)
        check("L+M basis detected from the item text", items[1]["basis"] == "L+M" and items[0]["basis"] == "L")
        # row 2: cleaner, 2 h.  row 3: explicit plumber (rate left blank).  row 4: cleaner + material.  row 5: no hours.
        data["sections"][0]["items"][0].update(hours=2)
        data["sections"][0]["items"][1].update(hours=1, trade="plumber")
        data["sections"][0]["items"][2].update(hours=1, needs_materials=True)
        sched.write_text(json.dumps(data), encoding="utf8")
        inputs = work / "inputs.json"
        inputs.write_text(json.dumps({"overhead": 0.10, "profit": 0.10, "tax": 0.0825, "rates": {"cleaner": 30}}), encoding="utf8")
        xlsx = work / "price.xlsx"
        cp = run(HERE / "price_workbook.py", "build", sched, xlsx, "--inputs", inputs)
        check("workbook builds", xlsx.exists(), cp.stdout + cp.stderr)
        import openpyxl
        wb = openpyxl.load_workbook(xlsx)
        check("workbook has the expected sheets",
              [ws.title for ws in wb.worksheets] == ["README", "Inputs", "Schedule A", "Attachment B", "Site Table", "Benchmarks", "Checks"],
              str([ws.title for ws in wb.worksheets]))
        check("inputs prefilled, reviewed gate left blank", wb["Inputs"]["B3"].value == 0.10 and wb["Inputs"]["B8"].value is None)
        soffice = shutil.which("soffice")
        if soffice:
            wb["Schedule A"]["H4"] = 20  # material cost for the cleaner L+M line (row 4)
            filled = work / "price_filled.xlsx"
            wb.save(filled)
            outdir = work / "calc"
            subprocess.run([soffice, f"-env:UserInstallation=file://{work}/lo", "--headless", "--convert-to",
                            "xlsx:Calc MS Excel 2007 XML", "--outdir", str(outdir), str(filled)], capture_output=True)
            calc = outdir / "price_filled.xlsx"
            if calc.exists():
                c = openpyxl.load_workbook(calc, data_only=True)["Schedule A"]
                check("labor-only line = 2h x $30 x 1.10 x 1.10 = 72.60", c["J2"].value == 72.6, repr(c["J2"].value))
                check("line whose trade rate is blank stays blank (never $0)", c["J3"].value in (None, ""), repr(c["J3"].value))
                expect = round((1 * 30 + 20 * 1.0825) * 1.21, 2)
                check(f"labor + taxed material line = {expect}", c["J4"].value == expect, repr(c["J4"].value))
                check("line with no hours stays blank", c["J5"].value in (None, ""), repr(c["J5"].value))
                ck = openpyxl.load_workbook(calc, data_only=True)["Checks"]
                ready = [ck.cell(r, 2).value for r in range(3, 25) if ck.cell(r, 1).value == "READY TO SUBMIT?"]
                check("checks sheet says NOT READY", ready == ["NOT READY"], str(ready))
            else:
                print("  skip  LibreOffice did not produce a recalculated file; formula checks skipped")
        else:
            print("  skip  soffice not installed; formula recalculation checks skipped")

        # shaded (agency-greyed) site-table cells: not inputs, not counted toward READY
        data["site_table"] = {"sizes": ["Efficiency", "1 Bdrm", "2 Bdrm"], "properties": [
            {"name": "P1", "units": 10, "address": "a", "shaded_cells": ["Efficiency", "2 Bdrm"]},
            {"name": "P2", "units": 5, "address": "b"}]}
        sched2 = work / "schedule_shaded.json"
        sched2.write_text(json.dumps(data), encoding="utf8")
        xlsx2 = work / "price_shaded.xlsx"
        run(HERE / "price_workbook.py", "build", sched2, xlsx2)
        wb2 = openpyxl.load_workbook(xlsx2)
        check("shaded cells are grey and not inputs; open cells are inputs",
              wb2["Site Table"]["E2"].fill.fgColor.rgb.endswith("D9D9D9") and wb2["Site Table"]["F2"].fill.fgColor.rgb.endswith("FFF8DC"))
        if soffice:
            wb2["Site Table"]["F2"] = 100   # P1, open
            wb2["Site Table"]["E2"] = 999   # P1, shaded: must NOT count
            wb2["Site Table"]["E3"] = 50    # P2, open
            f2 = work / "price_shaded_f.xlsx"
            wb2.save(f2)
            out2 = work / "calc2"
            subprocess.run([soffice, f"-env:UserInstallation=file://{work}/lo", "--headless", "--convert-to",
                            "xlsx:Calc MS Excel 2007 XML", "--outdir", str(out2), str(f2)], capture_output=True)
            if (out2 / "price_shaded_f.xlsx").exists():
                ck2 = openpyxl.load_workbook(out2 / "price_shaded_f.xlsx", data_only=True)["Checks"]
                rows = {ck2.cell(r, 1).value: ck2.cell(r, 2).value for r in range(3, 25) if ck2.cell(r, 1).value}
                exp = next(v for k, v in rows.items() if k.startswith("Site-table price cells expected"))
                fil = next(v for k, v in rows.items() if k.startswith("Site-table price cells filled"))
                check("site table expects only the 4 open cells and ignores a value typed into a grey one", (exp, fil) == (4, 2), f"{exp},{fil}")

        # -------------------------------------------------- compliance matrix
        print("compliance_matrix.py")
        reg = work / "reg"
        reg.mkdir()
        good = [{"id": "S1-001", "lines": "1", "section": "s", "type": "scope", "requirement": "Start in 24h",
                 "quote": "begin within 24 hours", "response_location": "Tab 3", "owner": "Contractor", "risk": "high", "note": ""}]
        (reg / "shred-1.json").write_text(json.dumps(good), encoding="utf8")
        (reg / "shred-1-anomalies.md").write_text("- term stated two ways\n", encoding="utf8")
        (reg / "shred-2-property-table.json").write_text(json.dumps({"not": "a register"}), encoding="utf8")
        cp = run(HERE / "compliance_matrix.py", reg, work / "mx", "--title", "T")
        check("clean register builds (helper json ignored)", cp.returncode == 0 and "1 rows from 1 files" in cp.stdout, cp.stdout)
        bad = good + [dict(good[0]), {"id": "S1-003", "type": "scope", "owner": "Contractor", "requirement": "", "risk": "extreme"}]
        (reg / "shred-3.json").write_text(json.dumps(bad), encoding="utf8")
        cp = run(HERE / "compliance_matrix.py", reg, work / "mx2", "--title", "T", expect_rc=None)
        check("duplicate id, missing field and bad risk are all reported",
              cp.returncode == 1 and "duplicate id S1-001" in cp.stdout and "missing requirement" in cp.stdout and "not low/medium/high" in cp.stdout,
              cp.stdout)
        cp = run(HERE / "compliance_matrix.py", "--verify", work / "mx.json", expect_rc=None)
        check("--verify exits 1 while rows are open", cp.returncode == 1)
        rows_now = json.loads((work / "mx.json").read_text(encoding="utf8"))
        rows_now[0]["status"], rows_now[0]["answered_in"], rows_now[0]["status_note"] = "answered", "Tab 3, section 1", "ok"
        (work / "mx.json").write_text(json.dumps(rows_now), encoding="utf8")
        cp = run(HERE / "compliance_matrix.py", "--rebuild", work / "mx.json", work / "mx", "--title", "T")
        kept = json.loads((work / "mx.json").read_text(encoding="utf8"))[0]
        check("--rebuild keeps entered statuses and notes (a rebuild from shreds would reset them to open)",
              cp.returncode == 0 and kept["status"] == "answered" and kept["status_note"] == "ok" and "'answered': 1" in cp.stdout, cp.stdout)
        cp = run(HERE / "compliance_matrix.py", "--verify", work / "mx.json", expect_rc=None)
        check("--verify exits 0 once every row has a status", cp.returncode == 0, cp.stdout)

        # --------------------------------------------- md -> docx + placeholders
        print("md_to_docx.py + scan_placeholders.py")
        md = work / "draft.md"
        md.write_text("---\ntitle: T\nheader_right: \"{{firm.legal_name}}\"\n---\n\n# Hello\n\nWe are {{firm.legal_name}}, est. {{firm.year_established}}.\n\n"
                      "| a | b |\n|---|---|\n| 1 | [[FILL: price]] |\n", encoding="utf8")
        intake = work / "intake.yaml"
        intake.write_text("firm:\n  legal_name: Example LLC\n  year_established: null\n", encoding="utf8")
        out_docx = work / "draft.docx"
        cp = run(HERE / "md_to_docx.py", md, out_docx, "--intake", intake)
        check("md_to_docx builds and reports one unresolved token", out_docx.exists() and "unresolved fact tokens=1" in cp.stdout, cp.stdout)
        cp = run(HERE / "scan_placeholders.py", out_docx, "--json", work / "scan.json", expect_rc=None)
        scan = json.loads((work / "scan.json").read_text(encoding="utf8"))[0]
        check("scanner finds exactly the 2 [[FILL]] flags in the docx", len(scan["fill"]) == 2 and len(scan["token"]) == 0, str(scan))
        check("scanner exits 1 while anything is unfinished", cp.returncode == 1)
        clean = work / "clean.md"
        clean.write_text("All done.\n", encoding="utf8")
        cp = run(HERE / "scan_placeholders.py", clean, expect_rc=None)
        check("scanner exits 0 on a finished file", cp.returncode == 0)
        wmd = work / "widths.md"
        wmd.write_text("| # | Why it matters | Pri |\n|---|---|---|\n| 1 | " + "A long explanation of why this matters to the price. " * 4 + " | P1 |\n\n"
                       "<!--\nmulti-line note to the writer\nmust not print\n-->\n\n"
                       "| Step {w=1} | What happens {w=5} |\n|---|---|\n| 1 | Notice arrives. |\n", encoding="utf8")
        wdocx = work / "widths.docx"
        run(HERE / "md_to_docx.py", wmd, wdocx)
        wd = docx.Document(wdocx)
        t1 = [round(c.width.inches, 2) for c in wd.tables[0].rows[0].cells]
        t2 = [round(c.width.inches, 2) for c in wd.tables[1].rows[0].cells]
        check("table columns are sized to their text (long column wider than a narrow one)", t1[1] > 3 * t1[0] and t1[0] == t1[2], str(t1))
        check("{w=N} header markers set relative widths and are removed from the text",
              abs(t2[1] / t2[0] - 5) < 0.2 and wd.tables[1].rows[0].cells[0].text == "Step", str(t2))
        check("multi-line HTML comments are dropped", not any("multi-line note" in p.text for p in wd.paragraphs))

        # ------------------------------------------------------ build_package
        print("build_package.py")
        pk = work / "pk"
        (pk / "src").mkdir(parents=True)
        for n in ("ext", "intl"):
            (pk / "src" / f"{n}.md").write_text("# Tab\n\nWe are {{firm.legal_name}}. [[FILL: x]]\n", encoding="utf8")
        (pk / "intake.yaml").write_text("firm:\n  legal_name: Example LLC\n", encoding="utf8")
        agency = None
        try:
            import pymupdf
            agency = pk / "agency.pdf"
            a = pymupdf.open()
            for i in range(3):
                a.new_page().insert_text((72, 72), f"agency page {i + 1}")
            a.save(agency)
        except ImportError:
            pass
        manifest = {"intake": "intake.yaml", "docs": [
            {"src": "src/ext.md", "name": "Tab1", "audience": "external"},
            {"src": "src/intl.md", "name": "Memo", "audience": "internal"}]}
        if agency:
            manifest["docs"][0]["append_pdf_pages"] = [{"pdf": "agency.pdf", "pages": [1, 3], "title": "Agency pages"}]
        (pk / "package.yaml").write_text(json.dumps(manifest), encoding="utf8")  # JSON is valid YAML
        cp = run(HERE / "build_package.py", pk / "package.yaml", "--no-pdf", expect_rc=None)
        check("build_package builds docx for both docs and exits 1 while flags remain",
              cp.returncode == 1 and (pk / "out/docx/Tab1.docx").exists() and (pk / "out/docx/Memo.docx").exists(), cp.stdout)
        check("build report counts 1 [[FILL]] per document", cp.stdout.count("| 1 | 0 | 0 |") == 2, cp.stdout)
        if soffice and agency:
            cp = run(HERE / "build_package.py", pk / "package.yaml", expect_rc=None)
            import pymupdf
            tab = pymupdf.open(pk / "out/pdf/Tab1.pdf")
            bundle = pymupdf.open(pk / "out/Proposal-Package.pdf")
            check("agency pages appended after a title page (doc + 1 title + 2 pages)", len(tab) >= 4 and "agency page 3" in tab[len(tab) - 1].get_text(), str(len(tab)))
            check("bundle holds external documents only, with one bookmark", len(bundle) == len(tab) and len(bundle.get_toc()) == 1, f"{len(bundle)} vs {len(tab)}")
        else:
            print("  skip  PDF build checks need soffice and PyMuPDF")
    finally:
        if not keep:
            shutil.rmtree(work, ignore_errors=True)
    print(f"\n{passed} checks passed, {len(failed)} failed")
    for f in failed:
        print("  FAILED:", f)
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
